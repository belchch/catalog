from __future__ import annotations

import docx


def extract_text(path: str, kind: str) -> str:
    """Return the plain text of a stored document.

    - ``md`` / ``result_md`` -> raw file contents (utf-8).
    - ``docx`` -> body paragraphs and tables in document order; tables as markdown.
    - ``csv`` -> raw file contents (utf-8 with cp1251/latin-1 fallback).
    - ``xlsx`` -> all non-empty sheets rendered as markdown tables.
    - ``pdf`` -> per-page text via pypdf with explicit page markers.
    """
    if kind in ("md", "result_md"):
        with open(path, encoding="utf-8") as f:
            return f.read()
    if kind == "docx":
        return _extract_docx(path)
    if kind == "csv":
        return _extract_csv(path)
    if kind == "xlsx":
        return _extract_xlsx(path)
    if kind == "pdf":
        return _extract_pdf(path)
    raise ValueError(f"unsupported kind: {kind}")


def _extract_docx(path: str) -> str:
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            parts.append(Paragraph(child, document).text)
        elif isinstance(child, CT_Tbl):
            parts.append(_docx_table_to_md(Table(child, document)))
    return "\n".join(parts)


def _docx_table_to_md(table: object) -> str:
    from docx.table import Table

    if not isinstance(table, Table):
        return ""
    rows = list(table.rows)
    if not rows:
        return ""
    width = max(len(row.cells) for row in rows)
    if width == 0:
        return ""
    padded: list[list[str]] = []
    nested_blocks: list[str] = []
    seen_cells: set[int] = set()
    for row in rows:
        values = [_cell_to_str(cell.text.strip()) for cell in row.cells]
        if len(values) < width:
            values = values + [""] * (width - len(values))
        padded.append(values)
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen_cells:
                continue
            seen_cells.add(tc_id)
            for nested in cell.tables:
                nested_md = _docx_table_to_md(nested)
                if nested_md:
                    nested_blocks.append(nested_md)
    header = padded[0]
    body = padded[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    result = "\n".join(lines)
    if nested_blocks:
        return result + "\n\n" + "\n\n".join(nested_blocks)
    return result


def _extract_csv(path: str) -> str:
    """Read a CSV file with utf-8, falling back to cp1251 then latin-1."""
    with open(path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _extract_xlsx(path: str) -> str:
    """Render all non-empty sheets of an xlsx workbook as markdown tables."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        blocks: list[str] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            rows = [list(r) for r in rows]
            while rows and all(_cell_is_blank(c) for c in rows[-1]):
                rows.pop()
            if not rows:
                continue
            block = [f"## Sheet: {sheet.title}"]
            width = max(len(r) for r in rows)
            padded = [
                [_cell_to_str(c) for c in r] + [""] * (width - len(r)) for r in rows
            ]
            header = padded[0]
            body = padded[1:]
            block.append("| " + " | ".join(header) + " |")
            block.append("| " + " | ".join(["---"] * width) + " |")
            for row in body:
                block.append("| " + " | ".join(row) + " |")
            blocks.append("\n".join(block))
        return "\n\n".join(blocks)
    finally:
        workbook.close()


def _extract_pdf(path: str) -> str:
    """Extract per-page text from a PDF with page markers.

    A scanned PDF (no text layer) yields an explicit warning instead of an
    empty string so callers can tell silent failures from real content.
    """
    from pypdf import PdfReader

    reader = PdfReader(path)
    page_texts: list[str] = []
    for page in reader.pages:
        page_texts.append((page.extract_text() or "").strip())

    if not any(page_texts):
        return (
            "[PDF has no extractable text layer (likely a scanned document); "
            "OCR is not supported in this build.]"
        )

    blocks: list[str] = []
    for index, text in enumerate(page_texts, start=1):
        blocks.append(f"\n\n--- page {index} ---\n\n{text}")
    return "\n".join(blocks).strip()


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    return text.replace("|", "\\|")


def _cell_is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
